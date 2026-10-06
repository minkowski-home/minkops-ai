"""Generate disposable mock-client bills and matching source data (no secrets)."""

import json
import sys
from pathlib import Path

import fitz
from openpyxl import Workbook
from openpyxl.worksheet.table import Table
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parents[1] / "apps/solution-api/data/min119-hard-demo"


def main():
    if "--refresh-scans" in sys.argv:
        for path in sorted(ROOT.glob("*.pdf")):
            if not path.stem.endswith("-scan"):
                render_scan(path)
        return
    if ROOT.exists():
        raise SystemExit("Demo already exists; keep its tested data or choose a fresh directory.")
    ROOT.mkdir(parents=True)
    bills = [
        ("01-new", "MOCK-101", 1000, 180),
        ("02-new", "MOCK-102", 2000, 360),
        ("03-backfill", "MOCK-099", 500, 90),
        ("04-correction", "MOCK-101", 1200, 216),
    ]
    for name, invoice, subtotal, tax in bills:
        path = ROOT / f"{name}.pdf"
        doc = canvas.Canvas(str(path))
        doc.setTitle(f"Synthetic mock-client bill {invoice}")
        # Dense, slightly skewed, low-contrast forms require visual reading and
        # business reasoning. All noise is authored in the PDF, not added by an
        # image editor; ground truth stays out of synchronized input folders.
        doc.saveState()
        doc.translate(12, 8)
        doc.rotate(-1.2)
        doc.setFillColorRGB(0.22, 0.25, 0.28)
        doc.setFont("Times-Bold", 16)
        doc.drawString(30, 785, "MDS MATERIALS / TAX INVOICE")
        doc.setFont("Helvetica", 7)
        doc.drawString(
            30, 770, "Registered entity: Minkops Demo Supplier | MOCK GSTIN: 29ABCDE1234F1Z5"
        )
        doc.drawString(30, 758, "Synthetic data for model evaluation. Not a valid tax document.")
        doc.rect(28, 655, 525, 92)
        doc.line(288, 655, 288, 747)
        doc.setFont("Helvetica-Bold", 9)
        doc.drawString(36, 733, "BILL TO: MOCK CLIENT")
        doc.drawString(298, 733, "SHIP TO / consignee: Site Stores")
        doc.setFont("Helvetica", 8)
        for x, y, value in [
            (36, 717, "Buyer GSTIN 29PQRSX4321K1Z8"),
            (36, 702, "PO: PO-9007 / dated 28-Sep-26"),
            (36, 687, "Bill period: Sep-Oct; financial year 26/27"),
            (298, 717, "GRN: GRN-00418 / receipt 02-Oct-26"),
            (298, 702, "Dispatch note: DN-211 / date 30-Sep-26"),
            (298, 687, "Project: MOCK-SITE; route reference QX-17"),
        ]:
            doc.drawString(x, y, value)
        doc.drawString(36, 667, f"Tax document no.: {invoice}")
        doc.drawString(298, 667, "Issued: 01 / OCT / 2026")
        doc.setFont("Helvetica-Bold", 8)
        headings = [
            (36, "Item / HSN"),
            (230, "Qty"),
            (295, "Rate (Rs)"),
            (380, "Gross"),
            (465, "Assessable"),
        ]
        for x, value in headings:
            doc.drawString(x, 625, value)
        doc.line(28, 616, 553, 616)
        doc.setFont("Times-Roman", 9)
        # Gross and discount deliberately differ from the taxable base.
        gross = subtotal + 40
        for y, values in [
            (
                596,
                [
                    "Aggregate blend / 2517",
                    "2.000 MT",
                    f"{gross / 2:,.2f}",
                    f"{gross:,.2f}",
                    f"{subtotal:,.2f}",
                ],
            ),
            (576, ["Less trade allowance", "", "", "(40.00)", "included above"]),
            (556, ["Returnable pallets", "4 NOS", "0.00", "0.00", "non-billable"]),
        ]:
            for (x, _), value in zip(headings, values, strict=True):
                doc.drawString(x, y, value)
        doc.line(28, 540, 553, 540)
        doc.setFont("Helvetica", 8)
        doc.drawString(36, 524, "Tax base after allowance - see allocation memo on reverse")
        labels = [
            ("Assessable value", subtotal),
            ("Central levy @ 9%", tax / 2),
            ("State levy @ 9%", tax / 2),
            ("Integrated levy", 0),
            ("Rounding adjustment", 0),
            ("AMOUNT PAYABLE", subtotal + tax),
        ]
        for index, (label, amount) in enumerate(labels):
            y = 494 - index * 22
            doc.setFont("Helvetica-Bold" if index == 5 else "Helvetica", 9)
            doc.drawString(308, y, label)
            doc.drawRightString(540, y, f"{amount:,.2f}")
        doc.setFont("Times-Italic", 8)
        doc.drawString(
            36, 370, "Bank remittance advice refers to the tax document number, not PO / GRN."
        )
        doc.drawString(
            36, 354, "Previous outstanding balance: Rs 7,080.00 - excluded from this invoice."
        )
        doc.drawString(
            36, 338, f"Duplicate office copy ref: {invoice}; this is one invoice, not another bill."
        )
        doc.setStrokeColorRGB(0.82, 0.82, 0.8)
        # Crossing form lines and a skewed receipt stamp make OCR ordering hard.
        for offset in range(8):
            doc.line(25, 270 + offset * 9, 552, 280 + offset * 9)
        doc.saveState()
        doc.translate(125, 440)
        doc.rotate(18)
        doc.setFillColorRGB(0.58, 0.38, 0.38)
        doc.setFont("Courier-Bold", 13)
        doc.drawString(0, 0, "RECEIVED 02 OCT 2026")
        doc.restoreState()
        doc.setFont("Helvetica", 7)
        doc.drawString(
            36,
            82,
            "Page 1 of 2 - amounts on the reverse repeat this document; do not double count.",
        )
        doc.restoreState()
        doc.showPage()
        doc.setFont("Helvetica-Bold", 14)
        doc.drawString(42, 780, "ALLOCATION MEMO / NOT AN ADDITIONAL INVOICE")
        doc.setFont("Helvetica", 9)
        memo = [
            f"Applies only to tax document {invoice}, issued 01 October 2026.",
            "Counterparty ledger: Minkops Demo Supplier (trading name MDS Materials).",
            "Purchase allocation: Minkops Demo Purchases; assessed base after Rs 40 allowance.",
            "Both central and state levy components post to Minkops Demo Tax in this mock company.",
            "The customer is Mock Client, not the supplier. Delivery date is not the invoice date.",
            "Project reference MOCK-SITE is for Excel reporting. No Tally cost allocation in this demo.",
            f"Control reconciliation: base {subtotal:.2f} + combined levy {tax:.2f} = due {subtotal + tax:.2f}.",
            "PO / GRN / delivery-note identifiers are supporting references, not bill identity.",
            "This mock accounting allocation does not represent a real GST filing configuration.",
        ]
        for index, line in enumerate(memo):
            doc.drawString(42, 735 - index * 30, line)
        doc.save()
        render_scan(path)
    (ROOT / "05-unreadable.png").write_bytes((ROOT / "01-new.png").read_bytes())
    # A valid blank image, rather than corrupt upload bytes: extraction must flag it.
    from PIL import Image

    Image.new("RGB", (600, 800), "white").save(ROOT / "05-unreadable.png")
    book = Workbook()
    suppliers = book.active
    suppliers.title = "Suppliers"
    suppliers.append(["Supplier", "Ledger"])
    suppliers.append(["Minkops Demo Supplier", "Minkops Demo Supplier"])
    projects = book.create_sheet("Projects")
    projects.append(["Project"])
    projects.append(["MOCK-SITE"])
    register = book.create_sheet("Bills")
    register.append(["Mock purchase register"])
    register.append(
        [
            "Row.Token",
            "Doc.Ref",
            "Party.External",
            "Posting.Day",
            "Allocation.Ref",
            "Txn.Base",
            "Levy.Sum",
            "Settlement.Due",
        ]
    )
    register.append(
        ["SEED-100", "MOCK-100", "Minkops Demo Supplier", "2026-10-01", "MOCK-SITE", 100, 18, 118]
    )
    register.add_table(Table(displayName="Purchases", ref="A2:H3"))
    book.save(ROOT / "mock-register.xlsx")
    expected = ROOT.parent / "min119-hard-ground-truth.json"
    expected.write_text(
        json.dumps(
            {
                "bills": bills,
                "mapping": {
                    "Doc.Ref": "invoice_number",
                    "Party.External": "vendor",
                    "Posting.Day": "date",
                    "Allocation.Ref": "cost_code",
                    "Txn.Base": "subtotal",
                    "Levy.Sum": "tax",
                    "Settlement.Due": "total",
                },
            },
            indent=2,
        )
    )
    print(ROOT)


def render_scan(path):
    with fitz.open(path) as pdf, fitz.open() as scanned:
        for index, page in enumerate(pdf):
            pix = page.get_pixmap(dpi=105)
            if index == 0:
                pix.save(path.with_suffix(".png"))
            target = scanned.new_page(width=page.rect.width, height=page.rect.height)
            target.insert_image(target.rect, stream=pix.tobytes("png"))
        scanned.save(path.with_name(path.stem + "-scan.pdf"), deflate=True)


if __name__ == "__main__":
    main()
