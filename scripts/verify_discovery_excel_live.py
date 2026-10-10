"""Opt-in MIN-124 paid verification of an Excel-authored named-table fixture.

Use the dedicated local database and a separate MINKOPS_LIVE_DATA checkpoint
directory. This installs the core Bill Entry definition only in that local DB;
the mock-client composition is restored afterwards. No production changes.
"""
import base64
import copy
import io
import json
import os
import uuid
import zipfile

import psycopg
from openpyxl import load_workbook
from minkops_platform.installation import install, load_solution
from minkops_platform.workflows import load_definition, register_workflow

import verify_discovery_context_live as proof

ROOT = os.environ.get("MINKOPS_EXCEL_FIXTURE_ROOT", proof.NATIVE_ROOT + r"\apps\solution-api\data\min124")
FILE = ROOT + r"\register.xlsx"


def main():
    proof.connect()
    api, state = proof.api, proof.state
    projected = proof.native({"operation": "schema", "path": FILE})
    content = base64.b64decode(projected["content"])
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        assert all(b"PRIVATE-DISCOVERY-CANARY" not in archive.read(p)
                   and b"PRIVATE-VENDOR-CANARY" not in archive.read(p) for p in archive.namelist())
    book = load_workbook(io.BytesIO(content))
    completed = state.get("result", {}).get("state") == "completed"
    assert book["Bills"].tables["MIN124Bills"].ref == ("A4:F6" if completed else "A4:F5")
    for row in range(5, 9):
        assert book["Bills"][f"G{row}"].value == f"=D{row}+E{row}"
    if "source_id" not in state:
        source = api("post", proof.base + "/accounts/sources", headers=proof.csrf,
            data={"paths": '["register.xlsx"]', "writable": "true"},
            files=[("files", ("register.xlsx", content))])
        state["source_id"] = source["id"]
        api("post", proof.base + f"/desktop/devices/{state['device_id']}/sources/{source['id']}", headers=proof.csrf)
        proof.checkpoint()
    # Workbook roles are operator configuration, not accounting data entry.
    # Explicitly review this fixture's intended destination before confirmation.
    def review_mapping(result):
        reviewed = copy.deepcopy(result)
        for sheet in reviewed['sheets']:
            if sheet['sheet'] == 'Bills' and sheet['table'] == 'MIN124Bills':
                sheet['role'] = 'destination'
        return reviewed
    scan = proof.discovery({"depth": "client_context", "destination_mode": "excel", "excel_source_ids": [state["source_id"]], "tally": None}, ROOT, review_mapping)
    mapping = scan["mapping_run"]["result"]["sheets"][0]
    assert mapping["table"] == "MIN124Bills" and mapping["header_row"] == 4
    assert mapping["role"] == "destination"
    with psycopg.connect(proof.DATABASE) as c:
        tenant = c.execute("SELECT id FROM tenants WHERE slug='mock-tenant'").fetchone()[0]
        register_workflow(c, tenant, load_definition(proof.ROOT / "employees/accounts-desk/workflows/bill-entry"))
    try:
        if "run_id" not in state:
            api("post", proof.base + "/desktop/jobs", headers=proof.csrf, json={
                "device_id": state["device_id"], "request_key": str(uuid.uuid4()),
                "operation": "files.refresh", "input": {"source_id": state["source_id"]}})
            proof.finish_job("refresh", ROOT)
            import pymupdf
            doc = pymupdf.open()
            page = doc.new_page()
            state["invoice"] = "MIN124-EXCEL-" + uuid.uuid4().hex[:8]
            page.insert_text((60, 80), "\n".join([
                "SYNTHETIC TEST INVOICE", "Invoice Number: " + state["invoice"],
                "Vendor: MIN124 Excel Supplier", "Invoice Date: 01 October 2026",
                "Subtotal: INR 300.00", "Tax: INR 0.00", "Total: INR 300.00",
            ]), fontsize=12)
            source = api("post", proof.base + "/accounts/sources", headers=proof.csrf,
                data={"paths": '["excel-bill.pdf"]'}, files=[("files", ("excel-bill.pdf", doc.tobytes()))])
            doc.close()
            run = api("post", proof.base + "/accounts/runs", headers=proof.csrf, json={
                "key": "bill-entry", "request_key": str(uuid.uuid4()),
                "file_ids": [source["files"][0]["id"]],
                "catalog_id": scan["catalog"]["excel_catalog_id"],
                "config": {"output_mode": "excel_in_place", "discovery_id": scan["id"]}})
            state["run_id"] = run["id"]
            proof.checkpoint()
        path = proof.base + f"/accounts/runs/{state['run_id']}"
        run = api("get", path)
        if run["state"] == "collecting":
            proof.finish_job("refresh", ROOT)
        proof.hosted(state["run_id"])
        run = api("get", path)
        if run["state"] == "review":
            assert not run["result"]["unresolved"] and len(run["result"]["records"]) == 1
            record = run["result"]["records"][0]
            assert record["table"] == "MIN124Bills"
            expected = {"Invoice Number": state["invoice"], "Vendor": "MIN124 Excel Supplier",
                        "Invoice Date": "2026-10-01", "Subtotal": 300, "Tax": 0, "Total": 300}
            assert all(record["data"].get(k) == v for k, v in expected.items()), record
            run = api("post", path + "/approve", headers=proof.csrf,
                      json={"result": run["result"], "acknowledge_findings": True})
        while run["state"] == "writing":
            proof.finish_job("save_excel", ROOT)
            run = api("get", path)
        assert run["state"] == "completed", run
        saved = proof.native({"operation": "read_file", "path": FILE})
        sheet = load_workbook(io.BytesIO(base64.b64decode(saved["content"])), data_only=False)["Bills"]
        assert sheet.tables["MIN124Bills"].ref == "A4:F6"
        assert sheet["A5"].value == "PRIVATE-DISCOVERY-CANARY"
        assert sheet["B5"].value == "PRIVATE-VENDOR-CANARY"
        assert sheet["D5"].value == 10 and sheet["E5"].value == 0 and sheet["F5"].value == 10
        matches = [row for row in sheet.iter_rows(min_row=5, max_col=6, values_only=True)
                   if row[0] == state["invoice"]]
        assert len(matches) == 1 and matches[0][1] == "MIN124 Excel Supplier" and matches[0][5] == 300
        for row in range(5, 9):
            assert sheet[f"G{row}"].value == f"=D{row}+E{row}"
        state["result"] = run
        proof.checkpoint()
        print("Excel structural-only discovery, paid extraction, reviewed native save and receipt replay passed.", flush=True)
    finally:
        with psycopg.connect(proof.DATABASE) as c:
            install(c, load_solution(proof.ROOT, "mock-client"), actor_email="demo@example.com")


if __name__ == "__main__":
    main()
