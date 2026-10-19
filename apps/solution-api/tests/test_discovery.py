"""Discovery uses real database dispatch and confirmation; no fake local executor."""

import base64
import hashlib
import os
import unittest
import uuid
from io import BytesIO

import psycopg
import test_desktop
from psycopg.rows import dict_row
from test_accounts_excel import catalog, workbook

URL = os.getenv("TEST_DATABASE_URL")


@unittest.skipUnless(URL, "TEST_DATABASE_URL is required")
class DiscoveryTests(unittest.TestCase):
    setUp = test_desktop.DesktopTests.setUp
    pair = test_desktop.DesktopTests.pair

    def test_native_only_collection_does_not_apply_hosted_interpretation_limits(self):
        from unittest.mock import patch

        run, _ = self.start(ai_context=False)
        with patch('minkops_platform.accounts.client_context.context_assets',
                   side_effect=ValueError('Complete collection exceeds hosted interpretation limits.')) as assets:
            current = self.collect(run)
        assets.assert_not_called()
        self.assertFalse(current['catalog']['partial'])
        confirmed = self.client.post(self.discovery + f"/runs/{run['id']}/confirm",
                                     headers=self.csrf, json={})
        self.assertEqual(confirmed.status_code, 200, confirmed.text)
        self.assertTrue(confirmed.json()['ready'])

    def test_hosted_collection_still_holds_an_oversized_interpretation(self):
        from unittest.mock import patch

        run, _ = self.start(ai_context=True)
        with patch('minkops_platform.accounts.client_context.context_assets',
                   side_effect=ValueError('Complete collection exceeds hosted interpretation limits.')):
            current = self.collect(run, analyze=False)
        self.assertTrue(current['catalog']['partial'])
        self.assertIsNone(current['mapping_run'])
        self.assertEqual(self.client.post(self.discovery + f"/runs/{run['id']}/confirm",
                                         headers=self.csrf, json={}).status_code, 409)

    def test_tally_without_ai_goes_directly_to_review_and_confirms_full_catalog(self):
        run, _ = self.start(ai_context=False)
        current = self.collect(run)
        self.assertEqual(current["state"], "review")
        self.assertIsNone(current["mapping_run"])
        self.assertFalse(current["ready"])
        self.assertEqual(current["catalog"]["context_notes"], [])
        confirmed = self.client.post(self.discovery + f"/runs/{current['id']}/confirm",
                                     headers=self.csrf, json={})
        self.assertEqual(confirmed.status_code, 200, confirmed.text)
        self.assertTrue(confirmed.json()["ready"])
        with psycopg.connect(URL, row_factory=dict_row) as c:
            saved = c.execute("SELECT catalog,mapping_run_id FROM discovery_runs WHERE id=%s",
                              (current["id"],)).fetchone()
        self.assertIsNone(saved["mapping_run_id"])
        # New catalogs deliver compact records; the stored export stays complete.
        self.assertEqual(saved["catalog"]["context_record_format"], "1")
        company = saved["catalog"]["sources"][0]["snapshot"]["companies"][0]
        self.assertEqual(company["masters"], {"status": "ready", "count": 0, "records": []})
        self.assertEqual(company["vouchers"], {"status": "ready", "count": 0, "records": []})
        self.assertNotIn("master_summary", company)

    def test_partial_tally_without_ai_cannot_be_confirmed(self):
        run, _ = self.start(ai_context=False)
        current = self.collect(run, partial=True)
        self.assertIsNone(current["mapping_run"])
        self.assertTrue(current["catalog"]["partial"])
        self.assertEqual(self.client.post(self.discovery + f"/runs/{run['id']}/confirm",
                                         headers=self.csrf, json={}).status_code, 409)

    def test_excel_without_ai_saves_structure_without_creating_entry_mappings(self):
        from minkops_connectors.excel import schema_projection
        from openpyxl import Workbook, load_workbook

        initial, _ = self.start()
        self.client.post(f"/api/tenants/{self.slug}/tasks/{initial['task_id']}/control",
                         headers=self.csrf, json={"action": "stop"}).raise_for_status()
        book = Workbook()
        book.active.append(["Vendor", "Total"])
        book.active["C2"] = "=SUM(B2:B10)"
        output = BytesIO()
        book.save(output)
        content, structure = schema_projection(output.getvalue())
        source = self.client.post(f"/api/tenants/{self.slug}/accounts/sources", headers=self.csrf,
                                  data={"paths": '["register.xlsx"]', "writable": "true"},
                                  files=[("files", ("register.xlsx", content))]).json()
        self.client.post(self.base + f"/devices/{self.device['id']}/sources/{source['id']}",
                         headers=self.csrf).raise_for_status()
        response = self.client.post(self.discovery + "/runs", headers=self.csrf, json={
            "device_id": self.device["id"], "request_key": str(uuid.uuid4()),
            "config": {"depth": "client_context", "ai_context": False,
                       "excel_source_ids": [source["id"]], "tally": None},
        })
        self.assertEqual(response.status_code, 202, response.text)
        run = response.json()
        claim = self.client.post(self.worker_base + "/claim", headers=self.worker, json={}).json()
        self.client.post(self.worker_base + f"/jobs/{claim['id']}/finish", headers=self.worker, json={
            "claim_token": claim["claim_token"], "result": {"sources": [{
                "key": source["id"], "tool": "excel", "status": "ready",
                "files": [{"path": "register.xlsx", "content": base64.b64encode(content).decode()}],
                "workbooks": [{"path": "register.xlsx", "status": "ready", "structure": structure,
                               "sha256": hashlib.sha256(content).hexdigest()}],
            }]},
        }).raise_for_status()
        current = self.client.get(self.discovery + f"/runs/{run['id']}").json()
        self.assertEqual(current["state"], "review")
        self.assertIsNone(current["mapping_run"])
        self.assertIsNone(current["catalog"]["excel_mappings"])
        confirmed = self.client.post(self.discovery + f"/runs/{run['id']}/confirm",
                                     headers=self.csrf, json={}).json()
        self.assertTrue(confirmed["ready"])
        self.assertNotIn("excel_catalog_id", confirmed["catalog"])
        file_id = confirmed["catalog"]["sources"][0]["workbooks"][0]["file_id"]
        with psycopg.connect(URL) as c:
            saved = c.execute("SELECT content FROM account_files WHERE id=%s", (file_id,)).fetchone()[0]
        self.assertEqual(load_workbook(BytesIO(bytes(saved))).active["C2"].value, "=SUM(B2:B10)")

    def test_explicit_source_choice_must_match_selected_tools(self):
        _, body = self.start()
        body["request_key"] = str(uuid.uuid4())
        body["config"]["destination_mode"] = "tally"
        accepted = self.client.post(self.discovery + "/runs", headers=self.csrf, json=body)
        self.assertEqual(accepted.status_code, 202, accepted.text)
        for mode in ("excel", "both"):
            with self.subTest(mode=mode):
                body["request_key"] = str(uuid.uuid4())
                body["config"]["destination_mode"] = mode
                rejected = self.client.post(self.discovery + "/runs", headers=self.csrf, json=body)
                self.assertEqual(rejected.status_code, 422, rejected.text)

    def test_paused_excel_tool_is_rejected_server_side(self):
        from minkops_platform.resources import REPOSITORY_ROOT
        from minkops_platform.solution_policy import bind_definition
        from minkops_platform.workflows import load_definition
        from psycopg.types.json import Jsonb

        _, body = self.start()
        bound = bind_definition(load_definition(
            REPOSITORY_ROOT / "employees/accounts-desk/workflows/source-discovery"), "mock-client")
        with psycopg.connect(URL) as c:
            c.execute("UPDATE workflows SET config_schema=%s WHERE key='source-discovery' AND tenant_id=(SELECT id FROM tenants WHERE slug=%s)",
                      (Jsonb(bound.tenant_schema), self.slug))
        source = self.client.post(f"/api/tenants/{self.slug}/accounts/sources", headers=self.csrf,
                                  data={"paths": '["register.xlsx"]', "writable": "true"},
                                  files=[("files", ("register.xlsx", workbook()))]).json()
        self.client.post(self.base + f"/devices/{self.device['id']}/sources/{source['id']}",
                         headers=self.csrf).raise_for_status()
        for tally in (None, body["config"]["tally"]):
            with self.subTest(tally=bool(tally)):
                rejected = self.client.post(self.discovery + "/runs", headers=self.csrf, json={
                    **body, "request_key": str(uuid.uuid4()),
                    "config": {**body["config"], "excel_source_ids": [source["id"]], "tally": tally}})
                self.assertEqual(rejected.status_code, 422, rejected.text)
                self.assertIn("does not use the selected source system", rejected.text)
        accepted = self.client.post(self.discovery + "/runs", headers=self.csrf,
                                    json={**body, "request_key": str(uuid.uuid4())})
        self.assertEqual(accepted.status_code, 202, accepted.text)

    def test_damaged_workbook_is_a_validation_error_not_a_backend_crash(self):
        r = self.client.post(
            f"/api/tenants/{self.slug}/accounts/sources",
            headers=self.csrf,
            data={"label": "Damaged", "writable": "true", "paths": '["bad.xlsx"]'},
            files=[("files", ("bad.xlsx", b"not an xlsx archive"))],
        )
        self.assertEqual(r.status_code, 422, r.text)

    def test_schema_registration_cannot_disguise_records_and_discovery_rejects_record_depth(self):
        rejected = self.client.post(
            f"/api/tenants/{self.slug}/accounts/sources",
            headers=self.csrf,
            data={"paths": '["books.xlsx"]', "schema_only": "true"},
            files=[("files", ("books.xlsx", workbook()))],
        )
        self.assertEqual(rejected.status_code, 422, rejected.text)
        _, body = self.start()
        body["request_key"] = str(uuid.uuid4())
        body["config"]["depth"] = "reference_data"
        rejected = self.client.post(self.discovery + "/runs", headers=self.csrf, json=body)
        self.assertEqual(rejected.status_code, 422, rejected.text)

    def test_excel_and_tally_collection_review_reuse_and_layout_change(self):
        from minkops_platform.accounts.worker import process
        from minkops_platform.resources import REPOSITORY_ROOT
        from minkops_platform.workflows import load_definition, register_workflow
        from openpyxl import load_workbook

        self.pair()
        self.discovery = f"/api/tenants/{self.slug}/discovery"
        accounts = f"/api/tenants/{self.slug}/accounts"
        definition = load_definition(
            REPOSITORY_ROOT / "employees/accounts-desk/workflows/source-discovery"
        )
        with psycopg.connect(URL) as c:
            tenant_id = c.execute("SELECT id FROM tenants WHERE slug=%s", (self.slug,)).fetchone()[
                0
            ]
            c.execute(
                "INSERT INTO employees(tenant_id,key,name) VALUES (%s,'accounts-desk','Accounts desk')",
                (tenant_id,),
            )
            workflow_id = register_workflow(c, tenant_id, definition)
            c.execute("UPDATE workflows SET status='active' WHERE id=%s", (workflow_id,))
        data = workbook()
        source = self.client.post(
            accounts + "/sources",
            headers=self.csrf,
            data={
                "label": "Mock books",
                "writable": "true",
                "paths": '["register.xlsx","bill.pdf"]',
            },
            files=[
                ("files", ("register.xlsx", data)),
                ("files", ("bill.pdf", b"%PDF-1.4 fixture")),
            ],
        ).json()
        self.assertEqual(
            self.client.post(
                self.base + f"/devices/{self.device['id']}/sources/{source['id']}",
                headers=self.csrf,
            ).status_code,
            200,
        )
        config = {"depth": "client_context", "excel_source_ids": [source["id"]], "tally": None}

        def scan(content):
            # Native discovery sends a new header-only package. Bill Entry must
            # refresh the original source independently before reading rows.
            projected = load_workbook(BytesIO(content))
            for sheet in projected:
                for row in sheet:
                    for cell in row:
                        if cell.row != (2 if sheet.title == "Bills" else 1):
                            cell.value = None
            projected_bytes = BytesIO()
            projected.save(projected_bytes)
            content = projected_bytes.getvalue()
            response = self.client.post(
                self.discovery + "/runs",
                headers=self.csrf,
                json={
                    "device_id": self.device["id"],
                    "request_key": str(uuid.uuid4()),
                    "config": config,
                },
            )
            self.assertEqual(response.status_code, 202, response.text)
            run = response.json()
            claim = self.client.post(
                self.worker_base + "/claim", headers=self.worker, json={}
            ).json()
            plan = self.client.get(
                self.worker_base + f"/jobs/{claim['id']}/plan?claim_token={claim['claim_token']}",
                headers=self.worker,
            )
            self.assertEqual(plan.status_code, 200, plan.text)
            for status in ("reading", "ready"):
                observed = self.client.post(
                    self.worker_base + f"/jobs/{claim['id']}/progress",
                    headers=self.worker,
                    json={
                        "claim_token": claim["claim_token"],
                        "source_key": source["id"],
                        "status": status,
                    },
                )
                self.assertEqual(observed.status_code, 200, observed.text)
            w = load_workbook(BytesIO(content))
            receipt = {
                "claim_token": claim["claim_token"],
                "result": {
                    "sources": [
                        {
                            "key": source["id"],
                            "tool": "excel",
                            "status": "ready",
                            "files": [
                                {
                                    "path": "register.xlsx",
                                    "content": base64.b64encode(content).decode(),
                                }
                            ],
                            "workbooks": [
                                {
                                    "path": "register.xlsx",
                                    "sha256": hashlib.sha256(content).hexdigest(),
                                    "status": "ready",
                                    "structure": {
                                        "sheets": [
                                            {
                                                "sheet": s.title,
                                                "tables": [],
                                                "preview": [
                                                    {
                                                        "row": 2 if s.title == "Bills" else 1,
                                                        "values": [
                                                            c.value
                                                            for c in s[
                                                                2 if s.title == "Bills" else 1
                                                            ]
                                                        ],
                                                    }
                                                ],
                                                "formulas": [],
                                                "reference_rows": [],
                                            }
                                            for s in w
                                        ]
                                    },
                                }
                            ],
                        }
                    ]
                },
            }
            for _ in range(2):
                r = self.client.post(
                    self.worker_base + f"/jobs/{claim['id']}/finish",
                    headers=self.worker,
                    json=receipt,
                )
                self.assertEqual(r.status_code, 200, r.text)
            return self.client.get(self.discovery + f"/runs/{run['id']}").json()

        first = scan(data)
        self.assertEqual(first["state"], "queued")
        mapping = first["mapping_run"]
        proposal = catalog()
        proposal["sheets"][0]["file_id"] = mapping["file_ids"][0]
        proposal["sheets"].append(
            {
                "file_id": mapping["file_ids"][0],
                "sheet": "Other",
                "table": None,
                "header_row": 1,
                "role": "ignore",
                "key_columns": [],
                "columns": [
                    {"name": "Keep me", "type": "string", "concept": "", "required": False}
                ],
            }
        )

        def executor(key, files, context, event, **kwargs):
            self.assertEqual(context["local_discovery"]["format_version"], "2")
            return proposal

        with psycopg.connect(URL, row_factory=dict_row) as c:
            process(
                c,
                c.execute("SELECT * FROM account_runs WHERE id=%s", (mapping["id"],)).fetchone(),
                executor=executor,
            )
        confirmed = self.client.post(
            self.discovery + f"/runs/{first['id']}/confirm",
            headers=self.csrf,
            json={"excel_mappings": proposal},
        )
        self.assertEqual(confirmed.status_code, 200, confirmed.text)
        self.assertTrue(confirmed.json()["ready"])
        w = load_workbook(BytesIO(data))
        w["Bills"].append(["A-2", "Acme", 24, "=C4*2"])
        buf = BytesIO()
        w.save(buf)
        second = scan(buf.getvalue())
        self.assertEqual(second["state"], "completed")
        self.assertTrue(second["ready"])
        self.assertEqual(
            second["catalog"]["excel_mappings"]["sheets"][0]["columns"],
            proposal["sheets"][0]["columns"],
        )
        self.assertIn(
            "bill.pdf",
            [f["path"] for f in self.client.get(accounts + "/sources").json()[0]["files"]],
        )
        w["Bills"]["C2"] = "Gross amount"
        buf = BytesIO()
        w.save(buf)
        changed = scan(buf.getvalue())
        self.assertEqual(changed["state"], "queued")
        self.assertFalse(changed["ready"])

    def start(self, ai_context=None, **extra):
        self.pair()
        from minkops_platform.resources import REPOSITORY_ROOT
        from minkops_platform.workflows import load_definition, register_workflow

        with psycopg.connect(URL) as c:
            tenant_id = c.execute("SELECT id FROM tenants WHERE slug=%s", (self.slug,)).fetchone()[
                0
            ]
            c.execute(
                "INSERT INTO employees(tenant_id,key,name) VALUES (%s,'accounts-desk','Accounts desk') ON CONFLICT DO NOTHING",
                (tenant_id,),
            )
            workflow = register_workflow(
                c,
                tenant_id,
                load_definition(
                    REPOSITORY_ROOT / "employees/accounts-desk/workflows/source-discovery"
                ),
            )
            c.execute("UPDATE workflows SET status='active' WHERE id=%s", (workflow,))
        self.discovery = f"/api/tenants/{self.slug}/discovery"
        body = {
            "device_id": self.device["id"],
            "request_key": str(uuid.uuid4()),
            "config": {
                "depth": "client_context",
                "excel_source_ids": [],
                "tally": {"port": 9000, "period": {"from": "2026-10-01", "to": "2026-10-02"}},
            },
        }
        if ai_context is not None:
            body["config"]["ai_context"] = ai_context
        body.update(extra)
        r = self.client.post(self.discovery + "/runs", headers=self.csrf, json=body)
        self.assertEqual(r.status_code, 202, r.text)
        return r.json(), body

    def collect(self, run, partial=False, *, vouchers=(), masters=(), number="100000", analyze=True):
        claim = self.client.post(self.worker_base + "/claim", headers=self.worker, json={}).json()
        result = {
            "sources": [
                {
                    "key": "tally",
                    "tool": "tally",
                    "status": "ready",
                    "snapshot": {
                        "format_version": "2",
                        "period": {"from": "2026-10-01", "to": "2026-10-02"},
                        "company_scope": [
                            {
                                "company": "Test",
                                "company_guid": "test-guid",
                                "company_number": number,
                            }
                        ],
                        "companies": [
                            {
                                "company": "Test",
                                "port": 9000,
                                "status": "unavailable",
                                "error": "Unavailable",
                            }
                        ]
                        if partial
                        else [
                            {
                                "company": "Test",
                                "port": 9000,
                                "status": "ready",
                                "company_guid": "test-guid",
                                "identity": {
                                    "NAME": "Test",
                                    "GUID": "test-guid",
                                    "COMPANYNUMBER": number,
                                },
                                "masters": {"status": "ready", "count": len(masters), "records": list(masters)},
                                "vouchers": {"status": "ready", "count": len(vouchers), "records": list(vouchers)},
                            }
                        ],
                        "partial": partial,
                    },
                }
            ]
        }
        r = self.client.post(
            self.worker_base + f"/jobs/{claim['id']}/finish",
            headers=self.worker,
            json={"claim_token": claim["claim_token"], "result": result},
        )
        self.assertEqual(r.status_code, 200, r.text)
        current = self.client.get(self.discovery + f"/runs/{run['id']}").json()
        if current["mapping_run"] and analyze:
            from minkops_platform.accounts.worker import process

            with psycopg.connect(URL, row_factory=dict_row) as c:
                process(
                    c,
                    c.execute(
                        "SELECT * FROM account_runs WHERE id=%s", (current["mapping_run"]["id"],)
                    ).fetchone(),
                    executor=lambda *args, **kwargs: {"sheets": [], "context_notes": []},
                )
        return self.client.get(self.discovery + f"/runs/{run['id']}").json()

    def test_partial_results_are_observable_downloadable_and_cannot_be_confirmed(self):
        run, _ = self.start()
        done = self.collect(run, partial=True)
        self.assertEqual(done["state"], "review")
        self.assertFalse(done["ready"])
        self.assertTrue(done["catalog"]["partial"])
        self.assertEqual(
            self.client.post(
                self.discovery + f"/runs/{run['id']}/confirm", headers=self.csrf, json={}
            ).status_code,
            409,
        )
        download = self.client.get(self.discovery + f"/runs/{run['id']}/catalog.json")
        self.assertEqual(download.status_code, 200)
        self.assertIn("attachment", download.headers["content-disposition"])
        self.assertEqual(
            download.json()["sources"][0]["snapshot"]["companies"][0]["status"], "unavailable"
        )

    def test_polling_returns_count_summaries_without_records_but_download_keeps_evidence(self):
        run, _ = self.start()
        self.collect(run)
        current = self.client.get(self.discovery + f"/runs/{run['id']}").json()
        company = current["catalog"]["sources"][0]["snapshot"]["companies"][0]
        self.assertIn("master_summary", company)
        self.assertNotIn("records", company["masters"])
        self.assertNotIn("records", company["vouchers"])
        self.assertNotIn("local_discovery_snapshot", current["mapping_run"]["config"])
        download = self.client.get(self.discovery + f"/runs/{run['id']}/catalog.json").json()
        self.assertIn("records", download["sources"][0]["snapshot"]["companies"][0]["masters"])

    def test_confirmation_and_unchanged_refresh_reuse_review_with_idempotence(self):
        run, body = self.start()
        self.assertEqual(
            self.client.post(self.discovery + "/runs", headers=self.csrf, json=body).json()["id"],
            run["id"],
        )
        self.collect(run)
        with psycopg.connect(URL, row_factory=dict_row) as connection:
            native_sources = connection.execute(
                "SELECT catalog->'sources' AS sources FROM discovery_runs WHERE id=%s",
                (run["id"],),
            ).fetchone()["sources"]
        confirmed = self.client.post(
            self.discovery + f"/runs/{run['id']}/confirm", headers=self.csrf, json={}
        )
        self.assertEqual(confirmed.status_code, 200, confirmed.text)
        self.assertTrue(confirmed.json()["ready"])
        self.assertIn(
            "master_summary",
            confirmed.json()["catalog"]["sources"][0]["snapshot"]["companies"][0],
        )
        with psycopg.connect(URL, row_factory=dict_row) as connection:
            stored_sources = connection.execute(
                "SELECT catalog->'sources' AS sources FROM discovery_runs WHERE id=%s",
                (run["id"],),
            ).fetchone()["sources"]
        self.assertEqual(stored_sources, native_sources)
        body["request_key"] = str(uuid.uuid4())
        refresh = self.client.post(self.discovery + "/runs", headers=self.csrf, json=body).json()
        again = self.collect(refresh)
        self.assertEqual(again["state"], "completed")
        self.assertTrue(again["ready"])
        self.assertEqual(again["catalog"]["review"]["reused_from"], run["id"])
        from minkops_platform.discovery import require_ready

        with psycopg.connect(URL, row_factory=dict_row) as c:
            tenant = c.execute("SELECT id FROM tenants WHERE slug=%s", (self.slug,)).fetchone()
            self.assertEqual(
                require_ready(c, tenant["id"], run["id"], ["tally"])["run_id"], run["id"]
            )

    def test_routine_tally_changes_keep_confirmation_but_company_changes_need_review(self):
        from minkops_platform.discovery import require_ready

        run, body = self.start(ai_context=False)
        self.collect(run)
        confirmed = self.client.post(self.discovery + f"/runs/{run['id']}/confirm", headers=self.csrf, json={})
        self.assertEqual(confirmed.status_code, 200, confirmed.text)
        # A voucher Minkops saved and a ledger created in Tally are routine data.
        voucher = {"type": "VOUCHER", "data": {"GUID": "saved-voucher", "DATE": "20261001"}}
        ledger = {"type": "LEDGER", "data": {"NAME": "New supplier", "GUID": "new-ledger",
                                            "PARENT": "Sundry Creditors"}}
        body["request_key"] = str(uuid.uuid4())
        refresh = self.client.post(self.discovery + "/runs", headers=self.csrf, json=body).json()
        again = self.collect(refresh, vouchers=[voucher], masters=[ledger])
        self.assertEqual((again["state"], again["ready"]), ("completed", True))
        self.assertEqual(again["catalog"]["review"]["reused_from"], run["id"])
        with psycopg.connect(URL, row_factory=dict_row) as c:
            tenant = c.execute("SELECT id FROM tenants WHERE slug=%s", (self.slug,)).fetchone()
            self.assertEqual(require_ready(c, tenant["id"], run["id"], ["tally"])["run_id"], run["id"])
        body["request_key"] = str(uuid.uuid4())
        changed = self.client.post(self.discovery + "/runs", headers=self.csrf, json=body).json()
        reviewed = self.collect(changed, number="100009")
        self.assertEqual((reviewed["state"], reviewed["ready"]), ("review", False))

    def test_failed_optional_analysis_leaves_a_complete_tally_collection_confirmable(self):
        from minkops_platform.accounts.run_store import AccountsRunStore

        run, _ = self.start()
        current = self.collect(run, analyze=False)
        with psycopg.connect(URL, row_factory=dict_row) as c:
            mapping = c.execute("SELECT * FROM account_runs WHERE id=%s",
                                (current["mapping_run"]["id"],)).fetchone()
            AccountsRunStore().fail(c, mapping, "Hosted workflow failed (APIError).")
        failed = self.client.get(self.discovery + f"/runs/{run['id']}").json()
        self.assertEqual((failed["state"], failed["analysis_failed"]), ("review", True))
        confirmed = self.client.post(self.discovery + f"/runs/{run['id']}/confirm", headers=self.csrf, json={})
        self.assertEqual(confirmed.status_code, 200, confirmed.text)
        self.assertTrue(confirmed.json()["ready"])
        self.assertEqual(confirmed.json()["catalog"]["context_notes"], [])
        with psycopg.connect(URL, row_factory=dict_row) as c:
            item = c.execute("SELECT status FROM attention_items WHERE event_key=%s",
                             (f"run:{mapping['id']}",)).fetchone()
        self.assertEqual(item["status"], "done")

    def test_scope_receipts_and_configuration_cannot_escape_the_requested_sources(self):
        run, _ = self.start()
        claim = self.client.post(self.worker_base + "/claim", headers=self.worker, json={}).json()
        route = self.worker_base + f"/jobs/{claim['id']}/finish"
        bad = {
            "claim_token": claim["claim_token"],
            "result": {
                "sources": [{"key": "foreign", "tool": "excel", "status": "ready", "snapshot": {}}]
            },
        }
        self.assertEqual(self.client.post(route, headers=self.worker, json=bad).status_code, 422)
        self.assertEqual(
            self.client.get("/api/tenants/mock-tenant/discovery/runs/" + run["id"]).status_code, 403
        )
        config = {
            "depth": "client_context",
            "excel_source_ids": [],
            "tally": {"company": "Test", "port": 9000, "categories": ["arbitrary"]},
        }
        self.assertEqual(
            self.client.post(
                self.discovery + "/runs",
                headers=self.csrf,
                json={
                    "device_id": self.device["id"],
                    "request_key": str(uuid.uuid4()),
                    "config": config,
                },
            ).status_code,
            422,
        )
        config["tally"] = None
        config["excel_source_ids"] = ["not-a-uuid"]
        self.assertEqual(
            self.client.post(
                self.discovery + "/runs",
                headers=self.csrf,
                json={
                    "device_id": self.device["id"],
                    "request_key": str(uuid.uuid4()),
                    "config": config,
                },
            ).status_code,
            422,
        )

    def test_measured_progress_is_optional_bounded_and_never_authorizes_completion(self):
        run, _ = self.start()
        claim = self.client.post(self.worker_base + "/claim", headers=self.worker, json={}).json()
        route = self.worker_base + f"/jobs/{claim['id']}/progress"
        body = {"claim_token": claim["claim_token"], "source_key": "tally", "status": "reading", "category": "A"}
        self.assertEqual(self.client.post(route, headers=self.worker, json=body).status_code, 200)
        with psycopg.connect(URL) as connection:
            before = connection.execute("SELECT count(*) FROM task_events WHERE task_id=%s", (run["task_id"],)).fetchone()[0]
        for completed in (2, 7):
            body["progress"] = {"completed": completed, "total": 7, "activity": "Collected masters", "item": "A"}
            response = self.client.post(route, headers=self.worker, json=body)
            self.assertEqual(response.status_code, 200, response.text)
        observed = self.client.get(self.discovery + f"/runs/{run['id']}").json()
        self.assertEqual(observed["observations"][0]["progress"]["completed"], 7)
        self.assertEqual(observed["state"], "executing")
        self.assertFalse(observed["ready"])
        with psycopg.connect(URL) as connection:
            after = connection.execute("SELECT count(*) FROM task_events WHERE task_id=%s", (run["task_id"],)).fetchone()[0]
        self.assertEqual(before, after)
        body["progress"]["completed"] = 8
        self.assertEqual(self.client.post(route, headers=self.worker, json=body).status_code, 422)

    def test_revocation_stops_discovery_and_progress_cannot_change_a_replaced_attempt(self):
        run, _ = self.start()
        claim = self.client.post(self.worker_base + "/claim", headers=self.worker, json={}).json()
        route = self.worker_base + f"/jobs/{claim['id']}/progress"
        r = self.client.post(
            route,
            headers=self.worker,
            json={"claim_token": str(uuid.uuid4()), "source_key": "tally", "status": "ready"},
        )
        self.assertEqual(r.status_code, 409)
        self.client.delete(self.base + f"/devices/{self.device['id']}", headers=self.csrf)
        self.assertEqual(
            self.client.get(self.discovery + f"/runs/{run['id']}").json()["state"], "failed"
        )

    def test_dependent_gate_rejects_partial_unconfirmed_and_superseded_catalogs(self):
        from minkops_platform.discovery import require_ready
        from minkops_platform.errors import ServiceError

        run, body = self.start()
        self.collect(run)
        with psycopg.connect(URL, row_factory=dict_row) as c:
            tenant = c.execute("SELECT id FROM tenants WHERE slug=%s", (self.slug,)).fetchone()
            with self.assertRaises(ServiceError):
                require_ready(c, tenant["id"], run["id"], ["tally"])
        self.client.post(self.discovery + f"/runs/{run['id']}/confirm", headers=self.csrf, json={})
        with psycopg.connect(URL, row_factory=dict_row) as c:
            self.assertEqual(
                require_ready(c, tenant["id"], run["id"], ["tally"])["review"]["status"],
                "confirmed",
            )
            with self.assertRaises(ServiceError):
                require_ready(c, tenant["id"], run["id"], ["excel"])
        body["request_key"] = str(uuid.uuid4())
        self.client.post(self.discovery + "/runs", headers=self.csrf, json=body)
        with psycopg.connect(URL, row_factory=dict_row) as c, self.assertRaises(ServiceError) as raised:
            require_ready(c, tenant["id"], run["id"], ["tally"])
        # A refresh still waiting for its PC has nothing to review yet.
        self.assertIn("still collecting", str(raised.exception))
