"""Source Discovery's Excel mappings must name real headers and tables."""

import unittest
from io import BytesIO

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

from minkops_connectors.excel import inspect_workbook
from minkops_platform.accounts.catalog import validate_catalog


def workbook():
    w = Workbook()
    s = w.active
    s.title = "Bills"
    s.append(["Company records"])
    s.append(["Invoice", "Vendor", "Amount", "Computed"])
    s.append(["A-1", "Acme", 12, "=C3*2"])
    s["A3"].number_format = "@"
    w.create_sheet("Other")["A1"] = "Keep me"
    b = BytesIO()
    w.save(b)
    return b.getvalue()


def catalog():
    return {
        "sheets": [
            {
                "file_id": "f1",
                "sheet": "Bills",
                "header_row": 2,
                "role": "destination",
                "key_columns": ["Invoice", "Vendor"],
                "columns": [
                    {
                        "name": "Invoice",
                        "type": "string",
                        "concept": "invoice_number",
                        "required": True,
                    },
                    {"name": "Vendor", "type": "string", "concept": "vendor", "required": True},
                    {"name": "Amount", "type": "number", "concept": "total", "required": True},
                    {"name": "Computed", "type": "number", "concept": "", "required": False},
                ],
            }
        ]
    }


class ExcelTests(unittest.TestCase):
    def test_proposed_semantic_conflicts_are_reviewable_but_not_confirmable(self):
        candidate = catalog()
        candidate["sheets"][0]["columns"][1]["concept"] = "invoice_number"
        validate_catalog(candidate, {"f1": workbook()}, review_proposal=True)
        with self.assertRaises(ValueError):
            validate_catalog(candidate, {"f1": workbook()})

    def test_a_named_table_mapping_must_use_that_tables_own_header(self):
        w = Workbook()
        s = w.active
        s.title = "Bills"
        for row in [("Invoice", "Vendor", "Amount"), ("A-1", "Acme", 12)]:
            s.append(row)
        s.add_table(Table(displayName="Purchases", ref="A1:C2"))
        b = BytesIO()
        w.save(b)
        m = catalog()["sheets"][0]
        m.update(table="Purchases", header_row=1)
        m["columns"] = m["columns"][:3]
        validate_catalog({"sheets": [m]}, {"f1": b.getvalue()})
        self.assertEqual(inspect_workbook(b.getvalue())[0]["tables"][0]["name"], "Purchases")
        with self.assertRaises(ValueError):
            validate_catalog({"sheets": [{**m, "header_row": 2}]}, {"f1": b.getvalue()})

    def test_headers_must_match_real_workbook(self):
        c = catalog()
        c["sheets"][0]["columns"][0]["name"] = "Invented"
        with self.assertRaises(ValueError):
            validate_catalog(c, {"f1": workbook()})

    def test_discovery_preserves_non_first_header_rows(self):
        inventory = inspect_workbook(workbook())
        self.assertEqual(inventory[0]["rows"][1][0], "Invoice")
        validate_catalog(catalog(), {"f1": workbook()})
