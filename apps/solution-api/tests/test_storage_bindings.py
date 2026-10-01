"""Legacy storage entry points preserve their output after package extraction."""

import json
from pathlib import Path
from tempfile import TemporaryDirectory
from unittest import TestCase
from unittest.mock import patch

from minkops_api import excel_writer, json_writer
from openpyxl import load_workbook


class StorageBindingTests(TestCase):
    def test_watermark_append_preserves_existing_rows_and_other_sheets(self):
        with TemporaryDirectory() as folder:
            destination = Path(folder) / "excel" / "watermarks.xlsx"
            with patch.object(excel_writer, "EXCEL_FILE", destination):
                self.assertEqual(
                    excel_writer.save_watermarks_to_excel([{"date": "2026-10-01", "latitude": 12}]),
                    destination,
                )
                workbook = load_workbook(destination)
                workbook.create_sheet("Reference")["A1"] = "keep me"
                workbook.save(destination)
                excel_writer.save_watermarks_to_excel([{"date": "2026-10-02", "address": "Pune"}])
            workbook = load_workbook(destination)
            self.assertEqual(workbook["Reference"]["A1"].value, "keep me")
            self.assertEqual(
                list(workbook["Watermark Data"].values),
                [
                    tuple(excel_writer.COLUMNS),
                    ("2026-10-01", None, 12, None, None, None),
                    ("2026-10-02", None, None, None, "Pune", None),
                ],
            )

    def test_json_artifact_keeps_envelope_and_unicode(self):
        with TemporaryDirectory() as folder:
            destination = Path(folder) / "json"
            with patch.object(json_writer, "JSON_DIR", destination):
                path = json_writer.save_extracted_json("bill.png", {"address": "पुणे"})
            self.assertEqual(path.parent, destination)
            record = json.loads(path.read_text())
            self.assertEqual(
                set(record), {"record_id", "source_filename", "created_at", "extracted_data"}
            )
            self.assertEqual(record["source_filename"], "bill.png")
            self.assertEqual(record["extracted_data"], {"address": "पुणे"})
            self.assertIn("पुणे", path.read_text())
