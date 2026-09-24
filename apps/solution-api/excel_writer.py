from pathlib import Path
from openpyxl import Workbook, load_workbook

BASE_DIR = Path(__file__).resolve().parent
EXCEL_DIR = BASE_DIR / "data" / "excel"
EXCEL_FILE = EXCEL_DIR / "watermark_data.xlsx"


COLUMNS = [
    "date",
    "time",
    "latitude",
    "longitude",
    "address",
    "altitude"
]

def save_watermarks_to_excel(watermarks):
    EXCEL_DIR.mkdir(parents = True, exist_ok=True)

    if EXCEL_FILE.exists():
        workbook = load_workbook(EXCEL_FILE)
        sheet = workbook["Watermark Data"]

    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Watermark Data"

        sheet.append(COLUMNS)

    for watermark in watermarks:
        row = [
            watermark.get("date"),
            watermark.get("time"),
            watermark.get("latitude"),
            watermark.get("longitude"),
            watermark.get("address"),
            watermark.get("altitude"),
        ]

        sheet.append(row)

    workbook.save(EXCEL_FILE)

    return EXCEL_FILE