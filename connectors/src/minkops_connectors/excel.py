"""Workbooks Minkops can read safely: a bounded .xlsx without parts openpyxl cannot round-trip."""

from io import BytesIO
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook

MAX_CELLS = 250_000


def open_workbook(content):
    try:
        archive = ZipFile(BytesIO(content))
    except (BadZipFile, OSError) as error:
        raise ValueError("The selected workbook is not a valid .xlsx file.") from error
    with archive:
        if sum(item.file_size for item in archive.infolist()) > 80_000_000:
            raise ValueError("Workbook expands beyond the supported size.")
        # openpyxl cannot safely round-trip every Excel extension.
        if any(
            "/slicer" in n or "/pivot" in n or "/activeX" in n or "/externalLink" in n
            for n in archive.namelist()
        ):
            raise ValueError(
                "Pivot, slicer, ActiveX and external-link workbooks are not supported yet."
            )
    w = load_workbook(BytesIO(content), data_only=False)
    if sum(s.max_row * s.max_column for s in w) > MAX_CELLS:
        raise ValueError("Workbook exceeds the supported cell limit.")
    return w
