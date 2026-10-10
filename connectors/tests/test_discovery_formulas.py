from io import BytesIO
import pytest
from openpyxl import Workbook, load_workbook
from minkops_connectors.excel import schema_projection


def test_projection_preserves_formula_text_without_business_values():
    book = Workbook()
    sheet = book.active
    sheet.append(["Invoice", "Total", "Computed"])
    sheet["C2"] = "=B2*2"
    stream = BytesIO()
    book.save(stream)
    content, structure = schema_projection(
        stream.getvalue(),
        {
            "sheets": [
                {
                    "sheet": sheet.title,
                    "tables": [],
                    "preview": [{"row": 1, "values": ["Invoice", "Total", "Computed"]}],
                    "formulas": [{"cell": "C2", "formula": "B2*2"}],
                }
            ]
        },
    )
    assert structure["sheets"][0]["formulas"] == [{"cell": "C2", "formula": "B2*2"}]
    assert load_workbook(BytesIO(content)).active["C2"].value == "=B2*2"
    sheet["A2"] = "hidden business row"
    stream = BytesIO()
    book.save(stream)
    with pytest.raises(ValueError):
        schema_projection(stream.getvalue())
